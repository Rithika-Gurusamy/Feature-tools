import io
import re
from typing import Tuple
from pypdf import PdfReader
from docx import Document

class ResumeParserService:
    @staticmethod
    def extract_text(file_bytes: bytes, filename: str) -> str:
        filename_lower = filename.lower()
        if filename_lower.endswith(".pdf"):
            return ResumeParserService._extract_pdf(file_bytes)
        elif filename_lower.endswith(".docx"):
            return ResumeParserService._extract_docx(file_bytes)
        elif filename_lower.endswith(".txt"):
            return file_bytes.decode("utf-8", errors="ignore").strip()
        else:
            raise ValueError(f"Unsupported file format: {filename}. Supported formats are .pdf, .docx, and .txt.")

    @staticmethod
    def _extract_pdf(file_bytes: bytes) -> str:
        reader = PdfReader(io.BytesIO(file_bytes))
        text_parts = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
        return "\n".join(text_parts).strip()

    @staticmethod
    def _extract_docx(file_bytes: bytes) -> str:
        doc = Document(io.BytesIO(file_bytes))
        text_parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)
        return "\n".join(text_parts).strip()

    @staticmethod
    def detect_candidate_name(resume_text: str, filename: str) -> str:
        lines = [line.strip() for line in resume_text.splitlines() if line.strip()]
        for line in lines[:5]:
            # Filter out common resume headers
            if re.search(r"resume|curriculum|vitae|email|phone|contact|github|linkedin", line, re.IGNORECASE):
                continue
            # Look for 2 to 4 words mostly alphabetic
            words = line.split()
            if 2 <= len(words) <= 4 and all(re.match(r"^[A-Za-z\.\'-]+$", w) for w in words):
                return line.title()

        # Fallback to filename without extension
        clean_fn = re.sub(r"[\._\-]*(resume|cv)[\._\-]*", "", filename, flags=re.IGNORECASE)
        clean_fn = re.sub(r"\.[a-zA-Z0-9]+$", "", clean_fn).strip()
        if clean_fn:
            words = re.split(r"[\s_\-]+", clean_fn)
            return " ".join(w.capitalize() for w in words if w)

        return "Candidate"

resume_service = ResumeParserService()
